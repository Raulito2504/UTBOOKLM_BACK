from pathlib import Path
from tempfile import NamedTemporaryFile
import uuid

from sqlalchemy.ext.asyncio import AsyncSession

from src.core.config import get_settings
from src.models import Document, DocumentChunk, DocumentStatus, IngestionJob, User
from src.modules.dashboard import service as dashboard_service
from src.modules.documents import repository
from src.modules.documents.storage import DocumentStorage, validate_upload_file
from src.modules.rag_chat.chunking import TextChunk, chunk_text


class DocumentNotFoundError(Exception):
    pass


class EmptyDocumentError(Exception):
    pass


class UnsupportedDocumentError(Exception):
    pass


class DocumentParseError(Exception):
    pass


class DocumentLimitExceededError(Exception):
    def __init__(self, error_code: str, message: str) -> None:
        self.error_code = error_code
        self.message = message
        super().__init__(message)


async def upload_document(
    db: AsyncSession,
    *,
    current_user: User,
    filename: str,
    content_type: str | None,
    content: bytes,
    title: str | None = None,
    storage: DocumentStorage,
) -> Document:
    if not content:
        raise EmptyDocumentError

    upload_file = _UploadFileStub(filename=filename, content_type=content_type)
    try:
        validate_upload_file(upload_file, len(content))
    except ValueError as exc:
        raise UnsupportedDocumentError(str(exc)) from exc

    parse_path = _write_temp_document(content, filename)
    try:
        chunks, page_count = parse_document_text(parse_path)
    except DocumentLimitExceededError:
        raise
    except ValueError as exc:
        raise UnsupportedDocumentError(str(exc)) from exc
    except RuntimeError as exc:
        raise DocumentParseError(str(exc)) from exc
    finally:
        Path(parse_path).unlink(missing_ok=True)

    file_path = storage.save(content, filename)

    try:
        document = await repository.create_document(
            db,
            organization_id=current_user.organization_id,
            uploaded_by_user_id=current_user.id,
            title=_document_title(filename=filename, title=title),
            file_path=file_path,
            file_size_bytes=len(content),
            page_count=page_count,
            original_filename=filename,
            mime_type=content_type,
            storage_backend=storage.backend,
            status=DocumentStatus.READY,
        )
        for chunk in chunks:
            await repository.create_document_chunk(
                db,
                document_id=document.id,
                chunk_index=chunk.index,
                content=chunk.content,
                page_number=chunk.page_number,
            )
        await dashboard_service.record_activity(
            db,
            current_user=current_user,
            activity_type="document_uploaded",
            metadata_json={
                "document_id": str(document.id),
                "title": document.title,
                "chunk_count": len(chunks),
            },
        )
    except Exception:
        storage.delete(file_path)
        raise
    return document


async def list_documents(
    db: AsyncSession,
    *,
    current_user: User,
    limit: int = 20,
    offset: int = 0,
) -> list[Document]:
    return await repository.list_documents_by_organization(
        db,
        organization_id=current_user.organization_id,
        limit=limit,
        offset=offset,
    )


async def get_document(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
) -> Document:
    document = await repository.get_document_by_organization(
        db,
        document_id=document_id,
        organization_id=current_user.organization_id,
    )
    if document is None:
        raise DocumentNotFoundError
    return document


async def update_document(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
    title: str | None = None,
) -> Document:
    document = await get_document(db, current_user=current_user, document_id=document_id)
    return await repository.update_document(
        db,
        document=document,
        title=title.strip() if title is not None else None,
    )


async def list_chunks(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
    limit: int = 20,
    offset: int = 0,
) -> list[DocumentChunk]:
    document = await get_document(db, current_user=current_user, document_id=document_id)
    return await repository.list_document_chunks(
        db,
        document_id=document.id,
        limit=limit,
        offset=offset,
    )


async def count_chunks(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
) -> int:
    document = await get_document(db, current_user=current_user, document_id=document_id)
    return await repository.count_document_chunks(db, document_id=document.id)


async def create_chunk(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
    content: str,
    page_number: int | None = None,
    chunk_index: int | None = None,
    tokens: int | None = None,
) -> DocumentChunk:
    settings = get_settings()
    document = await get_document(db, current_user=current_user, document_id=document_id)
    chunk_count = await repository.count_document_chunks(db, document_id=document.id)
    if chunk_count >= settings.document_max_chunks:
        raise DocumentLimitExceededError(
            "document_chunk_limit_exceeded",
            f"Document cannot have more than {settings.document_max_chunks} chunks",
        )
    if chunk_index is None:
        chunk_index = await repository.get_next_document_chunk_index(
            db,
            document_id=document.id,
        )
    return await repository.create_document_chunk(
        db,
        document_id=document.id,
        chunk_index=chunk_index,
        content=content,
        page_number=page_number,
        tokens=tokens,
    )


async def create_ingestion_job(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
) -> IngestionJob:
    document = await get_document(db, current_user=current_user, document_id=document_id)
    return await repository.create_ingestion_job(
        db,
        document_id=document.id,
        original_filename=document.original_filename,
        mime_type=document.mime_type,
        storage_backend=document.storage_backend,
    )


async def list_ingestion_jobs(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
) -> list[IngestionJob]:
    document = await get_document(db, current_user=current_user, document_id=document_id)
    return await repository.list_ingestion_jobs(db, document_id=document.id)


async def delete_document(
    db: AsyncSession,
    *,
    current_user: User,
    document_id: uuid.UUID,
    storage: DocumentStorage,
) -> None:
    document = await get_document(db, current_user=current_user, document_id=document_id)
    await repository.delete_document(db, document=document)
    storage.delete(document.file_path)


def _write_temp_document(content: bytes, filename: str) -> str:
    extension = Path(filename).suffix.lower()
    with NamedTemporaryFile(delete=False, suffix=extension) as temp_file:
        temp_file.write(content)
        return temp_file.name


def _document_title(*, filename: str, title: str | None) -> str:
    candidate = (title or "").strip()
    if candidate:
        return candidate
    return Path(filename).stem.strip() or "Untitled document"


def parse_document_text(file_path: str) -> tuple[list[TextChunk], int]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        chunks, page_count = parse_pdf(file_path)
    elif suffix == ".pptx":
        chunks, page_count = parse_pptx(file_path)
    elif suffix in {".md", ".txt"}:
        chunks, page_count = parse_plain_text(file_path)
    else:
        raise ValueError("Unsupported document type")
    _ensure_chunk_limit(chunks)
    return chunks, page_count


def parse_plain_text(file_path: str) -> tuple[list[TextChunk], int]:
    settings = get_settings()
    content = Path(file_path).read_bytes()
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("utf-8", errors="replace")
    if len(text) > settings.document_max_text_chars:
        raise DocumentLimitExceededError(
            "text_document_too_large",
            f"Text document cannot exceed {settings.document_max_text_chars} characters",
        )
    return _chunk_text(text, page_number=1), 1


def parse_pdf(file_path: str) -> tuple[list[TextChunk], int]:
    settings = get_settings()
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("PDF parser is not installed") from exc

    chunks: list[TextChunk] = []
    with fitz.open(file_path) as document:
        if document.page_count > settings.document_max_pages:
            raise DocumentLimitExceededError(
                "document_page_limit_exceeded",
                f"PDF cannot exceed {settings.document_max_pages} pages",
            )
        for page_index, page in enumerate(document):
            page_chunks = _chunk_text(page.get_text(), page_number=page_index + 1)
            chunks.extend(
                TextChunk(
                    index=len(chunks) + chunk.index,
                    content=chunk.content,
                    page_number=chunk.page_number,
                )
                for chunk in page_chunks
            )
        return chunks, document.page_count


def parse_pptx(file_path: str) -> tuple[list[TextChunk], int]:
    settings = get_settings()
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX parser is not installed") from exc

    presentation = Presentation(file_path)
    slide_count = len(presentation.slides)
    if slide_count > settings.document_max_slides:
        raise DocumentLimitExceededError(
            "document_page_limit_exceeded",
            f"PPTX cannot exceed {settings.document_max_slides} slides",
        )
    chunks: list[TextChunk] = []
    for slide_index, slide in enumerate(presentation.slides):
        text_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                text_parts.append(shape.text)
        for chunk in _chunk_text("\n".join(text_parts), page_number=slide_index + 1):
            chunks.append(
                TextChunk(
                    index=len(chunks),
                    content=chunk.content,
                    page_number=chunk.page_number,
                ),
            )
    return chunks, slide_count


def _chunk_text(text: str, *, page_number: int | None = None) -> list[TextChunk]:
    settings = get_settings()
    return chunk_text(
        text,
        chunk_size=settings.document_chunk_size,
        overlap=settings.document_chunk_overlap,
        page_number=page_number,
    )


def _ensure_chunk_limit(chunks: list[TextChunk]) -> None:
    settings = get_settings()
    if len(chunks) > settings.document_max_chunks:
        raise DocumentLimitExceededError(
            "document_chunk_limit_exceeded",
            f"Document cannot exceed {settings.document_max_chunks} chunks",
        )


class _UploadFileStub:
    def __init__(self, *, filename: str, content_type: str | None) -> None:
        self.filename = filename
        self.content_type = content_type
